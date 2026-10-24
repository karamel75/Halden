"""Couche 6 · Base de données : un classeur Excel lu et réécrit par le moteur.

Chaque onglet est une table. Le moteur charge tout en mémoire au début d'un cycle,
modifie les lignes, puis réécrit le classeur avec la charte Halden (nuit + champagne).
"""
from __future__ import annotations

import datetime as dt
from pathlib import Path
from typing import Any, Callable, Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

NUIT, CHAMPAGNE, BRONZE, PERLE, ARDOISE = "0B1426", "C9A86A", "8A6A3A", "F4F5F7", "3A4A63"
FONT = "Montserrat"

# Ordre des onglets et colonnes de chaque table.
SCHEMA: dict[str, list[str]] = {
    "Parametres": ["cle", "valeur", "description"],
    "Pack_Pays": ["cle", "valeur", "unite", "source", "statut"],
    "Grille_Decision": ["tache_id", "tache", "agent", "V1_volume", "V2_temps", "V3_cash_risque",
                        "C1_ambiguite", "C2_consequence", "C3_integration", "C4_jugement",
                        "niveau", "phase", "exemple"],
    "Projets": ["projet_id", "nom", "promoteur", "localisation", "statut", "avancement_pct",
                "date_lancement", "date_livraison", "compte_escrow", "banque_escrow"],
    "Jalons": ["jalon_id", "projet_id", "libelle", "seuil_pct", "date_certificat", "certificat_ref", "statut"],
    "Unites": ["unite_id", "projet_id", "etage", "type", "surface_sqft", "prix_aed", "statut",
               "acheteur_id", "date_reservation", "date_livraison"],
    "Acheteurs": ["acheteur_id", "nom", "nationalite", "langue", "canal", "email", "telephone",
                  "kyc_passeport", "kyc_emirates_id", "kyc_justif_adresse"],
    "Contrats": ["contrat_id", "unite_id", "acheteur_id", "date_reservation", "spa_statut",
                 "date_spa", "oqood_statut", "date_oqood"],
    "Echeancier": ["echeance_id", "projet_id", "unite_id", "acheteur_id", "num", "libelle", "declencheur",
                   "pourcentage", "montant_aed", "date_appel", "date_limite", "montant_paye_aed",
                   "statut", "reference"],
    "Releve_Bancaire": ["ligne_id", "date_valeur", "montant_aed", "libelle", "donneur_ordre",
                        "fichier_source", "statut", "echeance_id", "score", "methode"],
    "Relances": ["relance_id", "date", "echeance_id", "acheteur_id", "niveau", "jours_retard",
                 "canal", "statut"],
    "Messages_Entrants": ["msg_id", "date", "acheteur_id", "canal", "texte", "statut",
                          "categorie", "langue_detectee"],
    "Messages_Sortants": ["sortant_id", "date", "acheteur_id", "canal", "langue", "objet", "texte",
                          "agent", "statut", "cle"],
    "Leads": ["lead_id", "date", "nom", "source", "nationalite", "budget_aed", "recherche",
              "dernier_contact", "interactions", "statut", "score", "unite_proposee"],
    "Reserves": ["reserve_id", "date", "projet_id", "unite_id", "description", "localisation",
                 "statut", "categorie", "gravite", "sous_traitant", "date_sla"],
    "Sous_Traitants": ["st_id", "nom", "specialite", "contact"],
    "Garantie": ["reclamation_id", "date", "unite_id", "acheteur_id", "description", "statut",
                 "categorie", "garantie_applicable", "couverte"],
    "File_Approbation": ["file_id", "date", "agent", "tache_id", "niveau", "acheteur_id", "unite_id",
                         "resume", "brouillon", "montant_aed", "raisons", "statut", "decide_par",
                         "date_decision", "commentaire", "action", "payload", "cle"],
    "Journal_Audit": ["n", "horodatage", "date_simulee", "acteur", "couche", "action", "objet",
                      "detail", "niveau", "hash_precedent", "hash"],
}

MONEY_COLS = {"montant_aed", "montant_paye_aed", "prix_aed", "budget_aed"}
DATE_COLS = {c for cols in SCHEMA.values() for c in cols if c.startswith("date")}
WIDE_COLS = {"texte": 60, "brouillon": 60, "resume": 48, "description": 44, "detail": 60,
             "libelle": 34, "exemple": 46, "tache": 38, "raisons": 40, "payload": 30, "objet": 28}

LISEZ_MOI = [
    ("HALDEN · BASE DE DONNÉES DU POC", "titre"),
    ("Every promise, carried through.", "tag"),
    ("", ""),
    ("Ce classeur est la base du POC : le moteur le lit au début de chaque cycle et le réécrit à la fin.", ""),
    ("Toutes les données sont fictives (noms, montants, IBAN, téléphones). Aucun lien avec des personnes réelles.", ""),
    ("", ""),
    ("Vous pouvez modifier les onglets à la main puis relancer un cycle :", "h"),
    ("• Messages_Entrants : ajoutez une ligne avec statut NOUVEAU → l'agent Relation acquéreurs la traite.", ""),
    ("• Jalons : passez un jalon en CERTIFIÉ avec une date → l'agent Collections prépare les appels de fonds.", ""),
    ("• Grille_Decision : changez les notes V/C → le niveau d'autonomie de l'agent change au prochain cycle.", ""),
    ("• Pack_Pays : délais de relance, garanties… (règles du pays, à faire valider par un avocat).", ""),
    ("", ""),
    ("Onglets de sortie (écrits par les agents) :", "h"),
    ("• File_Approbation : tout ce qui attend une décision humaine (statut EN ATTENTE).", ""),
    ("• Messages_Sortants : messages préparés. En mode shadow, rien n'est envoyé.", ""),
    ("• Journal_Audit : chaque action, chaînée par hash (toute modification a posteriori se voit).", ""),
    ("", ""),
    ("Règle d'or : l'IA prépare, les règles calculent, un humain engage le promoteur. Jamais de mouvement d'argent.", "h"),
]


def _norm(v: Any) -> Any:
    if isinstance(v, dt.datetime):
        return v.date() if (v.hour, v.minute, v.second) == (0, 0, 0) else v
    return v


class Database:
    """Accès simple aux onglets : liste de dictionnaires par table."""

    def __init__(self, path: Path | str):
        self.path = Path(path)
        self.tables: dict[str, list[dict]] = {name: [] for name in SCHEMA}

    # -- lecture / écriture -------------------------------------------------
    @classmethod
    def load(cls, path: Path | str) -> "Database":
        db = cls(path)
        wb = load_workbook(db.path)
        for name, cols in SCHEMA.items():
            if name not in wb.sheetnames:
                continue
            ws = wb[name]
            rows = ws.iter_rows(values_only=True)
            header = [h for h in next(rows, [])]
            for r in rows:
                if r is None or all(v is None for v in r):
                    continue
                rec = {c: None for c in cols}
                for h, v in zip(header, r):
                    if h in rec:
                        rec[h] = _norm(v)
                db.tables[name].append(rec)
        return db

    def save(self, logo_png: Path | None = None) -> None:
        wb = Workbook()
        wb.remove(wb.active)
        self._write_cover(wb, logo_png)
        thin = Side(style="thin", color="D5D8DE")
        for name, cols in SCHEMA.items():
            ws = wb.create_sheet(name)
            ws.sheet_properties.tabColor = CHAMPAGNE if name in {"File_Approbation", "Journal_Audit"} else NUIT
            ws.append(cols)
            for c in ws[1]:
                c.font = Font(name=FONT, bold=True, color="FFFFFF", size=9)
                c.fill = PatternFill("solid", fgColor=NUIT)
                c.alignment = Alignment(vertical="center")
                c.border = Border(bottom=Side(style="medium", color=CHAMPAGNE))
            ws.row_dimensions[1].height = 22
            for rec in self.tables[name]:
                ws.append([rec.get(c) for c in cols])
            for j, col in enumerate(cols, start=1):
                letter = get_column_letter(j)
                ws.column_dimensions[letter].width = WIDE_COLS.get(col, max(12, min(24, len(col) + 4)))
                for cell in ws[letter][1:]:
                    cell.font = Font(name=FONT, size=9, color=NUIT)
                    cell.border = Border(bottom=thin)
                    if col in MONEY_COLS:
                        cell.number_format = '#,##0.00'
                    elif col in DATE_COLS:
                        cell.number_format = "yyyy-mm-dd"
                    if col in WIDE_COLS:
                        cell.alignment = Alignment(wrap_text=False, vertical="top")
            ws.freeze_panes = "B2"
            if self.tables[name]:
                ws.auto_filter.ref = ws.dimensions
        self.path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(self.path)

    def _write_cover(self, wb: Workbook, logo_png: Path | None) -> None:
        ws = wb.create_sheet("Lisez_moi")
        ws.sheet_properties.tabColor = CHAMPAGNE
        ws.sheet_view.showGridLines = False
        ws.column_dimensions["A"].width = 4
        ws.column_dimensions["B"].width = 120
        for r in range(1, 40):
            for col in ("A", "B", "C"):
                ws[f"{col}{r}"].fill = PatternFill("solid", fgColor=NUIT)
        start = 2
        if logo_png and Path(logo_png).exists():
            try:
                img = XLImage(str(logo_png))   # nécessite Pillow ; sans lui, couverture sans logo
                ratio = img.height / img.width
                img.width, img.height = 420, int(420 * ratio)
                ws.add_image(img, "B2")
                start = 12
            except ImportError:
                pass
        for i, (txt, kind) in enumerate(LISEZ_MOI):
            cell = ws[f"B{start + i}"]
            cell.value = txt
            color, size, bold = "E6E8EC", 10, False
            if kind == "titre":
                color, size, bold = "FFFFFF", 16, True
            elif kind == "tag":
                color, size = CHAMPAGNE, 10
            elif kind == "h":
                color, bold = CHAMPAGNE, True
            cell.font = Font(name=FONT, size=size, bold=bold, color=color)

    # -- requêtes --------------------------------------------------------------
    def rows(self, table: str) -> list[dict]:
        return self.tables[table]

    def where(self, table: str, pred: Callable[[dict], bool] | None = None, **eq) -> list[dict]:
        out = []
        for r in self.tables[table]:
            if all(r.get(k) == v for k, v in eq.items()) and (pred is None or pred(r)):
                out.append(r)
        return out

    def get(self, table: str, **eq) -> dict | None:
        res = self.where(table, **eq)
        return res[0] if res else None

    def add(self, table: str, rec: dict) -> dict:
        full = {c: None for c in SCHEMA[table]}
        full.update({k: v for k, v in rec.items() if k in full})
        self.tables[table].append(full)
        return full

    def add_many(self, table: str, recs: Iterable[dict]) -> None:
        for r in recs:
            self.add(table, r)

    def next_id(self, table: str, prefix: str, width: int = 4) -> str:
        key = SCHEMA[table][0]
        n = 0
        for r in self.tables[table]:
            v = str(r.get(key) or "")
            if v.startswith(prefix):
                try:
                    n = max(n, int(v[len(prefix):]))
                except ValueError:
                    pass
        return f"{prefix}{n + 1:0{width}d}"

    def param(self, key: str, default: Any = None) -> Any:
        r = self.get("Parametres", cle=key)
        return r["valeur"] if r else default

    def set_param(self, key: str, value: Any) -> None:
        r = self.get("Parametres", cle=key)
        if r:
            r["valeur"] = value

    def pack(self, key: str, default: Any = None) -> Any:
        r = self.get("Pack_Pays", cle=key)
        return r["valeur"] if r else default
